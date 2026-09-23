import json
import unittest
from io import BytesIO
from unittest.mock import patch

from docx import Document
from pypdf import PdfReader, PdfWriter

from app import (
    AI_JOBS,
    CHAT_SESSIONS,
    ScreeningHandler,
    _create_ai_job_record,
    _run_ai_review_job,
    calculate_contract_score,
    chat_with_ai,
    create_chat_session,
    extract_contract_text,
    make_report_pdf,
    render_ai_job_page,
    render_chat_page,
    render_home,
    render_report,
    review_with_ai,
    validate_request_origin,
)
from pii_masking import mask_pii
from screening import screen_text


class ScreeningTests(unittest.TestCase):
    def test_missing_wages_emits_review_finding_with_article_21_basis(self):
        result = screen_text(
            "HỢP ĐỒNG LAO ĐỘNG\nCông việc: Kế toán\nThời hạn: 12 tháng"
        )

        finding = next(
            item
            for item in result["findings"]
            if item["id"] == "BLLD2019-ART21-WAGES"
        )

        self.assertEqual(finding["severity"], "review")
        self.assertEqual(
            finding["legal_basis"],
            "Điều 21 khoản 1 điểm đ Bộ luật Lao động 2019",
        )

    def test_work_without_workplace_emits_workplace_review_finding(self):
        result = screen_text("Hợp đồng lao động. Công việc: Kế toán.")

        self.assertTrue(
            any(item["id"] == "BLLD2019-ART21-WORKPLACE" for item in result["findings"])
        )

    def test_fixed_term_over_36_months_emits_review_finding(self):
        result = screen_text("Hợp đồng xác định thời hạn 48 tháng. Công việc: Kế toán.")

        self.assertTrue(
            any(item["id"] == "BLLD2019-ART20-FIXED-TERM" for item in result["findings"])
        )

    def test_probation_over_60_days_emits_review_finding(self):
        result = screen_text(
            "Thư mời làm việc. Vị trí Nhân viên vận hành bảo mật hệ thống. Thời gian thử việc 90 ngày."
        )

        finding = next(
            item
            for item in result["findings"]
            if item["id"] == "BLLD2019-ART25-PROBATION-DURATION"
        )
        self.assertIn("90 ngày", finding["evidence"])

    def test_ai_report_keeps_important_uncited_rule_engine_findings(self):
        result = screen_text(
            "Thư mời làm việc. Vị trí Nhân viên vận hành bảo mật hệ thống. Thời gian thử việc 90 ngày."
        )
        ai_review = {
            "summary": "Sơ bộ.",
            "findings": [{
                "title": "Điểm khác",
                "issue": "Cần đối chiếu.",
                "suggested_revision": "Làm rõ điều khoản.",
                "legal_basis_ids": [],
                "missing_facts": [],
            }],
            "clarifying_questions": [],
        }

        html = render_report(result, "offer.docx", ai_review=ai_review)

        self.assertIn("Rule engine bổ sung", html)
        self.assertIn("thử việc 90 ngày", html)
        self.assertEqual(calculate_contract_score(result, ai_review)["value"], 92)

    def test_result_does_not_contain_compliance_verdict(self):
        result = screen_text("Công việc: Kế toán")

        self.assertNotIn("compliant", str(result).lower())
        self.assertNotIn("hợp pháp", str(result).casefold())

    def test_multi_law_rule_packs_emit_domain_findings(self):
        text = (
            "Hợp đồng lao động. Lương gross 50.000.000 VNĐ. "
            "Người lao động xử lý dữ liệu cá nhân khách hàng, mã nguồn và tác phẩm thuộc công ty."
        )
        ids = {item["id"] for item in screen_text(text)["findings"]}

        self.assertIn("BHXH2024-COMPULSORY-WAGE-BASE", ids)
        self.assertIn("PIT2007-GROSS-NET-WITHHOLDING", ids)
        self.assertIn("PDPD2023-EMPLOYEE-CUSTOMER-DATA", ids)
        self.assertIn("IP2005-WORK-CREATED-ASSET", ids)

    def test_base_contract_rules_cover_common_contract_gaps(self):
        result = screen_text("Hợp đồng dịch vụ. Bên A thuê Bên B tư vấn triển khai hệ thống.")
        ids = {item["id"] for item in result["findings"]}

        self.assertIn("BASE-CONTRACT-PARTIES-AUTHORITY", ids)
        self.assertIn("BASE-CONTRACT-PAYMENT", ids)
        self.assertIn("BASE-CONTRACT-TERMINATION", ids)
        self.assertIn("BASE-CONTRACT-DISPUTE", ids)

    def test_base_contract_rules_cover_delivery_confidentiality_ip_liability_and_notice(self):
        result = screen_text(
            "Hợp đồng dịch vụ. Bên A thuê Bên B triển khai phần mềm, bàn giao dữ liệu và tài liệu nội bộ. Nếu vi phạm gây thiệt hại sẽ xử lý theo hợp đồng."
        )
        ids = {item["id"] for item in result["findings"]}

        self.assertIn("BASE-CONTRACT-DELIVERY-ACCEPTANCE", ids)
        self.assertIn("BASE-CONTRACT-CONFIDENTIALITY", ids)
        self.assertIn("BASE-CONTRACT-IP-OWNERSHIP", ids)
        self.assertIn("BASE-CONTRACT-LIABILITY-CAP", ids)
        self.assertIn("BASE-CONTRACT-NOTICES", ids)

    def test_base_contract_rules_cover_term_amendment_assignment_order_and_language(self):
        result = screen_text(
            "Hợp đồng dịch vụ có hiệu lực từ ngày ký, kèm phụ lục tiếng Anh. Bên B có thể chuyển nhượng nghĩa vụ nếu cần."
        )
        ids = {item["id"] for item in result["findings"]}

        self.assertIn("BASE-CONTRACT-EFFECTIVE-TERM", ids)
        self.assertIn("BASE-CONTRACT-AMENDMENT", ids)
        self.assertIn("BASE-CONTRACT-ASSIGNMENT", ids)
        self.assertIn("BASE-CONTRACT-DOCUMENT-ORDER", ids)
        self.assertIn("BASE-CONTRACT-LANGUAGE", ids)

    def test_rule_packs_have_practical_mvp_coverage(self):
        result = screen_text("Hợp đồng dịch vụ. Bên B cung cấp dịch vụ vận hành website cho Bên A.")
        versions = result["rule_pack_version"]

        self.assertIn("0.4.0-base-contract", versions)
        self.assertIn("0.4.0-service-contract", versions)
        self.assertIn("0.4.0-cooperation-contract", versions)
        self.assertIn("0.3.0-labor-contract", versions)

    def test_service_contract_rules_cover_scope_acceptance_and_sla(self):
        result = screen_text("Hợp đồng dịch vụ. Bên B cung cấp dịch vụ vận hành website cho Bên A.")
        ids = {item["id"] for item in result["findings"]}

        self.assertIn("SERVICE-CONTRACT-SCOPE-DELIVERABLES", ids)
        self.assertIn("SERVICE-CONTRACT-ACCEPTANCE", ids)
        self.assertIn("SERVICE-CONTRACT-SLA", ids)

    def test_service_contract_rules_cover_change_request_data_security_and_subcontracting(self):
        result = screen_text(
            "Hợp đồng dịch vụ vận hành website và xử lý dữ liệu khách hàng. Bên B được thuê nhân sự hỗ trợ."
        )
        ids = {item["id"] for item in result["findings"]}

        self.assertIn("SERVICE-CONTRACT-CHANGE-REQUEST", ids)
        self.assertIn("SERVICE-CONTRACT-DATA-SECURITY", ids)
        self.assertIn("SERVICE-CONTRACT-SUBCONTRACTING", ids)
        self.assertIn("SERVICE-CONTRACT-WARRANTY-SUPPORT", ids)

    def test_service_contract_rules_cover_reporting_customer_duties_credits_and_exit_handover(self):
        result = screen_text(
            "Hợp đồng dịch vụ vận hành website. Bên A cung cấp tài khoản và dữ liệu đầu vào. Bên B báo cáo hàng tháng; nếu lỗi lặp lại sẽ hoàn tiền và bàn giao khi kết thúc."
        )
        ids = {item["id"] for item in result["findings"]}

        self.assertIn("SERVICE-CONTRACT-CUSTOMER-RESPONSIBILITIES", ids)
        self.assertIn("SERVICE-CONTRACT-REPORTING", ids)
        self.assertIn("SERVICE-CONTRACT-SERVICE-CREDITS", ids)
        self.assertIn("SERVICE-CONTRACT-EXIT-HANDOVER", ids)

    def test_service_contract_does_not_emit_labor_required_rules(self):
        result = screen_text("Hợp đồng dịch vụ. Bên B cung cấp dịch vụ vận hành website cho Bên A.")
        ids = {item["id"] for item in result["findings"]}

        self.assertNotIn("BLLD2019-ART21-WAGES", ids)
        self.assertNotIn("BLLD2019-ART21-EMPLOYEE", ids)

    def test_service_contract_does_not_emit_cooperation_rules_for_customer_data(self):
        result = screen_text("Hợp đồng dịch vụ. Bên B xử lý dữ liệu khách hàng để vận hành website cho Bên A.")
        ids = {item["id"] for item in result["findings"]}

        self.assertFalse(any(rule_id.startswith("COOP-CONTRACT-") for rule_id in ids))

    def test_cooperation_contract_rules_cover_contribution_profit_and_exit(self):
        result = screen_text("Hợp đồng hợp tác kinh doanh giữa hai bên để phát triển sản phẩm.")
        ids = {item["id"] for item in result["findings"]}

        self.assertIn("COOP-CONTRACT-CONTRIBUTIONS", ids)
        self.assertIn("COOP-CONTRACT-PROFIT-SHARING", ids)
        self.assertIn("COOP-CONTRACT-EXIT", ids)

    def test_cooperation_contract_rules_cover_governance_exclusivity_and_ip(self):
        result = screen_text(
            "Hợp đồng hợp tác kinh doanh phát triển nền tảng. Hai bên cùng khai thác khách hàng và sản phẩm chung."
        )
        ids = {item["id"] for item in result["findings"]}

        self.assertIn("COOP-CONTRACT-GOVERNANCE", ids)
        self.assertIn("COOP-CONTRACT-EXCLUSIVITY", ids)
        self.assertIn("COOP-CONTRACT-CUSTOMER-DATA", ids)
        self.assertIn("COOP-CONTRACT-IP-BRAND", ids)

    def test_cooperation_contract_rules_cover_accounting_deadlock_confidentiality_and_non_solicit(self):
        result = screen_text(
            "Hợp đồng hợp tác kinh doanh cùng phát triển sản phẩm. Hai bên chia doanh thu, sử dụng nhân sự và thông tin khách hàng chung."
        )
        ids = {item["id"] for item in result["findings"]}

        self.assertIn("COOP-CONTRACT-ACCOUNTING-AUDIT", ids)
        self.assertIn("COOP-CONTRACT-DEADLOCK", ids)
        self.assertIn("COOP-CONTRACT-CONFIDENTIALITY", ids)
        self.assertIn("COOP-CONTRACT-NON-SOLICIT", ids)

    def test_labor_contract_rules_cover_probation_overtime_leave_and_confidentiality(self):
        result = screen_text(
            "Hợp đồng lao động. Người lao động làm kỹ sư phần mềm, xử lý mã nguồn và dữ liệu khách hàng."
        )
        ids = {item["id"] for item in result["findings"]}

        self.assertIn("BLLD2019-ART24-PROBATION", ids)
        self.assertIn("BLLD2019-OVERTIME", ids)
        self.assertIn("BLLD2019-LEAVE", ids)
        self.assertIn("BLLD2019-CONFIDENTIALITY-IP", ids)

    def test_labor_contract_rules_cover_remote_work_equipment_deductions_and_termination(self):
        result = screen_text(
            "Hợp đồng lao động. Người lao động làm việc từ xa bằng laptop công ty. Lương có thể bị khấu trừ nếu nghỉ việc không bàn giao."
        )
        ids = {item["id"] for item in result["findings"]}

        self.assertIn("BLLD2019-REMOTE-WORK", ids)
        self.assertIn("BLLD2019-EQUIPMENT-EXPENSES", ids)
        self.assertIn("BLLD2019-SALARY-DEDUCTIONS", ids)
        self.assertIn("BLLD2019-TERMINATION-HANDOVER", ids)


class AppTests(unittest.TestCase):
    def test_mask_pii_masks_bank_account_address_and_labeled_names(self):
        text = (
            "Bên A: Nguyễn Văn A, địa chỉ: 12 Nguyễn Huệ, phường Bến Nghé, quận 1, TP.HCM. "
            "Số tài khoản: 012345678901 tại Vietcombank."
        )

        result = mask_pii(text)

        self.assertNotIn("Nguyễn Văn A", result.masked_text)
        self.assertNotIn("12 Nguyễn Huệ", result.masked_text)
        self.assertNotIn("012345678901", result.masked_text)
        self.assertIn("PERSON_NAME", result.masked_text)
        self.assertIn("ADDRESS", result.masked_text)
        self.assertIn("BANK_ACCOUNT", result.masked_text)

    def test_report_escapes_untrusted_filename_and_keeps_preliminary_disclaimer(self):
        html = render_report(
            screen_text("Công việc: Kế toán"),
            '<img src=x onerror=alert(1)>.txt',
        )

        self.assertIn("&lt;img", html)
        self.assertNotIn("<img src=x", html)
        self.assertIn("Kết quả là sàng lọc sơ bộ", html)
        self.assertIn("tham khảo", html)
        self.assertIn("mọi quyết định vẫn là CON NGƯỜI", html)

    def test_contract_score_penalizes_rule_and_ai_findings(self):
        result = screen_text("Hợp đồng dịch vụ. Bên B cung cấp dịch vụ vận hành website cho Bên A.")
        score = calculate_contract_score(result, {"summary": "OK", "findings": [{}, {}]})

        self.assertLess(score["value"], 100)
        self.assertGreaterEqual(score["value"], 0)
        self.assertIn(score["level"], {"Tốt", "Tạm chấp nhận", "Cần rà soát kỹ", "Rủi ro cao"})
        self.assertIn("weighted risk", score["explanation"])

    def test_contract_score_weights_high_risk_rules_more_than_admin_rules(self):
        result = {
            "findings": [
                {"id": "BASE-CONTRACT-LIABILITY-CAP", "title": "Cần kiểm tra giới hạn trách nhiệm", "missing_facts": []},
                {"id": "BASE-CONTRACT-PAYMENT", "title": "Cần kiểm tra thanh toán", "missing_facts": []},
                {"id": "BASE-CONTRACT-LANGUAGE", "title": "Cần kiểm tra ngôn ngữ", "missing_facts": []},
            ]
        }

        score = calculate_contract_score(result)

        self.assertEqual(score["deduction"], 19)
        self.assertIn("weighted risk từ rule engine", score["explanation"])

    def test_contract_score_ignores_findings_satisfied_by_pii_placeholders(self):
        result = {
            "findings": [
                {
                    "id": "BASE-CONTRACT-PARTIES-AUTHORITY",
                    "title": "Cần kiểm tra chủ thể và thẩm quyền ký",
                    "missing_facts": ["Thông tin định danh các bên"],
                },
                {
                    "id": "BASE-CONTRACT-NOTICES",
                    "title": "Cần kiểm tra thông báo và đầu mối liên hệ",
                    "missing_facts": ["Kênh gửi thông báo"],
                },
                {
                    "id": "BASE-CONTRACT-PAYMENT",
                    "title": "Cần kiểm tra thanh toán",
                    "missing_facts": [],
                },
            ]
        }
        contract_text = "Bên A {{PERSON_NAME_1}}, địa chỉ {{ADDRESS_1}}, MST {{TAX_ID_1}}, email {{EMAIL_ADDRESS_1}}."

        score = calculate_contract_score(result, contract_text=contract_text)

        self.assertEqual(score["deduction"], 8)
        self.assertIn("bỏ qua 2 finding do dữ liệu PII đã được masking", score["explanation"])

    def test_contract_score_with_ai_uses_visible_ai_findings_not_hidden_rule_count(self):
        result = {
            "findings": [
                {"id": "BASE-CONTRACT-PAYMENT", "missing_facts": []},
                {"id": "BASE-CONTRACT-LIABILITY-CAP", "missing_facts": []},
                {"id": "BASE-CONTRACT-TERMINATION", "missing_facts": []},
                {"id": "BASE-CONTRACT-DISPUTE", "missing_facts": []},
                {"id": "BLLD2019-ART21-WAGES", "missing_facts": []},
                {"id": "BLLD2019-ART21-INSURANCE", "missing_facts": []},
                {"id": "PDPD2023-EMPLOYEE-CUSTOMER-DATA", "missing_facts": []},
                {"id": "BHXH2024-COMPULSORY-WAGE-BASE", "missing_facts": []},
                {"id": "PIT2007-GROSS-NET-WITHHOLDING", "missing_facts": []},
                {"id": "SERVICE-CONTRACT-DATA-SECURITY", "missing_facts": []},
                {"id": "COOP-CONTRACT-DEADLOCK", "missing_facts": []},
                {"id": "BASE-CONTRACT-IP-OWNERSHIP", "missing_facts": []},
                {"id": "BLLD2019-OVERTIME", "missing_facts": []},
            ]
        }
        ai_review = {
            "summary": "Sơ bộ.",
            "findings": [
                {"legal_basis_ids": ["BHXH2024-COMPULSORY-WAGE-BASE"]},
                {"legal_basis_ids": ["BHXH2024-COMPULSORY-WAGE-BASE", "BLLD2019-OVERTIME"]},
            ],
            "clarifying_questions": [{}, {}, {}, {}],
        }

        score = calculate_contract_score(result, ai_review)

        self.assertEqual(score["deduction"], 16)
        self.assertEqual(score["value"], 84)
        self.assertIn("2 AI finding hiển thị", score["explanation"])

    def test_report_shows_contract_score_and_human_decision_disclaimer(self):
        ai_review = {"summary": "OK", "findings": [], "clarifying_questions": []}
        html = render_report(screen_text("Công việc: Kế toán"), "contract.txt", ai_review=ai_review)

        self.assertIn("Điểm tổng quan", html)
        self.assertIn("/100", html)
        self.assertIn("Đánh giá nhanh", html)
        self.assertIn("tham khảo", html)
        self.assertIn("CON NGƯỜI", html)

    def test_report_lists_legal_sources_but_hides_rule_checks(self):
        html = render_report(
            screen_text("Hợp đồng lao động. Công việc: Kế toán."),
            "contract.txt",
        )

        self.assertIn("Căn cứ pháp lý đang dùng", html)
        self.assertIn("Bộ luật Lao động 2019", html)
        self.assertIn("Luật Bảo hiểm xã hội 2024", html)
        self.assertIn("Nguồn: CSDL văn bản pháp luật", html)
        self.assertNotIn("Nguồn chính thức</a>", html)
        self.assertNotIn("<h2>Rule checks</h2>", html)
        self.assertNotIn("BLLD2019-ART21-WAGES", html)

    def test_ai_consent_says_pii_is_masked_before_ai_transfer(self):
        home = render_home()
        self.assertIn("Thông tin cá nhân sẽ được che", home)
        self.assertIn("trước khi gửi đến AI", home)
        self.assertNotIn("Authorization key", home)

    def test_home_shows_loading_overlay_after_submit(self):
        home = render_home()
        self.assertIn("id='loading-overlay'", home)
        self.assertIn("Đang rà soát hợp đồng", home)
        self.assertIn("addEventListener('submit'", home)

    def test_primary_button_text_is_centered(self):
        home = render_home()
        self.assertIn("justify-content:center", home)
        self.assertIn("align-items:center", home)

    def test_healthz_returns_ok(self):
        sent: dict[str, object] = {}

        class Handler(ScreeningHandler):
            path = "/healthz"
            def send_response(self, code, message=None): sent["status"] = code
            def send_header(self, keyword, value): sent[keyword] = value
            def end_headers(self): pass

        handler = object.__new__(Handler)
        handler.wfile = BytesIO()
        handler.do_GET()

        self.assertEqual(sent["status"], 200)
        self.assertEqual(handler.wfile.getvalue(), b"ok")

    def test_cross_site_origin_is_rejected(self):
        with self.assertRaisesRegex(ValueError, "Origin"):
            validate_request_origin("https://attacker.example", "127.0.0.1:8000")

    def test_local_origin_and_cli_without_origin_are_accepted(self):
        validate_request_origin("http://127.0.0.1:8000", "127.0.0.1:8000")
        validate_request_origin("https://vn-legal-contract-review.onrender.com", "vn-legal-contract-review.onrender.com")
        validate_request_origin("", "127.0.0.1:8000")

    def test_extract_contract_text_reads_docx_paragraphs_and_tables(self):
        stream = BytesIO()
        document = Document()
        document.add_paragraph("Hợp đồng dịch vụ")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Bên A"
        table.cell(0, 1).text = "Bên B cung cấp dịch vụ vận hành website"
        document.save(stream)

        text = extract_contract_text("contract.docx", stream.getvalue())

        self.assertIn("Hợp đồng dịch vụ", text)
        self.assertIn("Bên B cung cấp dịch vụ vận hành website", text)

    def test_extract_contract_text_rejects_invalid_extension(self):
        with self.assertRaisesRegex(ValueError, "TXT, PDF hoặc DOCX"):
            extract_contract_text("contract.doc", b"not a doc")

    def test_extract_contract_text_rejects_oversized_upload(self):
        with self.assertRaisesRegex(ValueError, "2 MiB"):
            extract_contract_text("contract.txt", b"a" * (2 * 1024 * 1024 + 1))

    def test_extract_contract_text_rejects_pdf_without_text_layer(self):
        stream = BytesIO()
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        writer.write(stream)

        with self.assertRaisesRegex(ValueError, "lớp văn bản"):
            extract_contract_text("scan.pdf", stream.getvalue())

    def test_report_pdf_contains_report_text(self):
        result = screen_text("Hợp đồng dịch vụ. Bên B cung cấp dịch vụ vận hành website cho Bên A.")
        pdf = make_report_pdf(result, "contract.docx")

        self.assertTrue(pdf.startswith(b"%PDF"))
        reader = PdfReader(BytesIO(pdf))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
        fonts = [
            str(font.get_object().get("/BaseFont"))
            for page in reader.pages
            for font in ((page.get("/Resources") or {}).get("/Font") or {}).values()
        ]
        self.assertIn("Kết quả rà soát sơ bộ", text)
        self.assertIn("contract.docx", text)
        self.assertIn("Bộ luật Dân sự 2015", text)
        self.assertIn("Điểm tổng quan", text)
        self.assertIn("/100", text)
        self.assertIn("CON NGƯỜI", text)
        self.assertNotIn("■", text)
        self.assertTrue(any("NotoSans" in font for font in fonts))

    def test_report_includes_pdf_download_button(self):
        html = render_report(
            screen_text("Hợp đồng dịch vụ. Bên B cung cấp dịch vụ vận hành website cho Bên A."),
            "contract.docx",
        )

        self.assertIn("download='contract-review-report.pdf'", html)
        self.assertIn("data:application/pdf;base64,", html)
        self.assertIn("Tải báo cáo PDF", html)

    def test_screen_with_ai_consent_redirects_to_background_job(self):
        body = (
            b"--x\r\n"
            b"Content-Disposition: form-data; name=\"contract\"; filename=\"contract.txt\"\r\n"
            b"Content-Type: text/plain\r\n\r\n"
            b"Cong viec: Ke toan\r\n"
            b"--x\r\n"
            b"Content-Disposition: form-data; name=\"ai_consent\"\r\n\r\n"
            b"yes\r\n"
            b"--x--\r\n"
        )
        sent: dict[str, object] = {"headers": {}}

        class Handler(ScreeningHandler):
            def send_response(self, code, message=None): sent["status"] = code
            def send_header(self, keyword, value): sent["headers"][keyword] = value
            def end_headers(self): pass

        handler = object.__new__(Handler)
        handler.headers = {
            "Content-Length": str(len(body)),
            "Content-Type": "multipart/form-data; boundary=x",
            "Host": "127.0.0.1:8000",
        }
        handler.rfile = BytesIO(body)
        handler.wfile = BytesIO()

        with patch("app.start_ai_review_job", return_value="job123") as mocked_job:
            handler._handle_screen()

        self.assertEqual(sent["status"], 303)
        self.assertEqual(sent["headers"]["Location"], "/job?id=job123")
        mocked_job.assert_called_once()

    def test_ai_job_page_auto_refreshes_until_review_finishes(self):
        AI_JOBS.clear()
        result = screen_text("Công việc: Kế toán")
        job_id = _create_ai_job_record("contract.txt", "Công việc: Kế toán", result)

        html = render_ai_job_page(job_id)

        self.assertIn("Đang rà soát bằng AI", html)
        self.assertIn("http-equiv='refresh'", html)
        self.assertIn(f"/job?id={job_id}", html)
        self.assertIn("<aside>Trang tự kiểm tra lại mỗi 5 giây. Bạn có thể để tab này mở.</aside>", html)
        self.assertNotIn("Cloudflare không còn phải chờ", html)

    def test_ai_job_page_renders_report_when_review_finishes(self):
        AI_JOBS.clear()
        result = screen_text("Công việc: Kế toán")
        ai_review = {"summary": "OK", "findings": [], "clarifying_questions": []}
        job_id = _create_ai_job_record("contract.txt", "Công việc: Kế toán", result)
        AI_JOBS[job_id]["status"] = "done"
        AI_JOBS[job_id]["ai_review"] = ai_review
        AI_JOBS[job_id]["session_id"] = create_chat_session("contract.txt", "Công việc: Kế toán", result, ai_review)

        html = render_ai_job_page(job_id)

        self.assertIn("Kết quả rà soát sơ bộ", html)
        self.assertIn("Hỏi tiếp về hợp đồng", html)
        self.assertNotIn("http-equiv='refresh'", html)

    def test_ai_review_job_stores_error_without_blocking_report(self):
        AI_JOBS.clear()
        result = screen_text("Công việc: Kế toán")
        job_id = _create_ai_job_record("contract.txt", "Công việc: Kế toán", result)

        with patch("app.review_with_ai", side_effect=RuntimeError("Không kết nối được AI provider.")):
            _run_ai_review_job(job_id)

        self.assertEqual(AI_JOBS[job_id]["status"], "error")
        self.assertIn("Không kết nối", AI_JOBS[job_id]["error"])
        html = render_ai_job_page(job_id)
        self.assertIn("Kết quả rule engine", html)
        self.assertIn("Không kết nối", html)

    def test_ai_review_requires_explicit_consent(self):
        with self.assertRaisesRegex(PermissionError, "đồng ý"):
            review_with_ai(
                "Công việc: Kế toán",
                screen_text("Công việc: Kế toán"),
                consent=False,
            )

    @patch("app.urlopen", side_effect=TimeoutError)
    @patch.dict("os.environ", {"LEGAL_AI_API_KEY": "test-key"}, clear=False)
    def test_ai_review_timeout_returns_safe_error(self, _mocked_urlopen):
        with self.assertRaisesRegex(RuntimeError, "Không kết nối được AI provider"):
            review_with_ai(
                "Công việc: Kế toán",
                screen_text("Công việc: Kế toán"),
                consent=True,
            )

    @patch("app.urlopen")
    @patch.dict("os.environ", {"LEGAL_AI_API_KEY": "test-key"}, clear=False)
    def test_ai_review_uses_allowed_rule_pack_citations_only(self, mocked_urlopen):
        class Response:
            def __enter__(self):
                return self

            def __exit__(self, *_):
                return False

            def read(self):
                return json.dumps(
                    {
                        "choices": [
                            {
                                "message": {
                                    "content": json.dumps(
                                        {
                                            "summary": "The contract is compliant.",
                                            "findings": [
                                                {
                                                    "title": "Cần kiểm tra lương",
                                                    "issue": "Thiếu mức lương rõ ràng.",
                                                    "risk_level": "review",
                                                    "legal_basis_ids": [
                                                        "BLLD2019-ART21-WAGES",
                                                        "MADE-UP-CITATION",
                                                    ],
                                                    "missing_facts": ["Mức lương"],
                                                },
                                                {
                                                    "title": "Cần kiểm tra không có căn cứ",
                                                    "issue": "Không được giữ finding này.",
                                                    "risk_level": "review",
                                                    "legal_basis_ids": ["MADE-UP-CITATION"],
                                                    "missing_facts": [],
                                                },
                                            ],
                                        },
                                        ensure_ascii=False,
                                    )
                                }
                            }
                        ]
                    },
                    ensure_ascii=False,
                ).encode()

        mocked_urlopen.return_value = Response()
        result = review_with_ai(
            "Công việc: Kế toán",
            screen_text("Công việc: Kế toán"),
            consent=True,
        )

        self.assertEqual(
            result["findings"][0]["legal_basis_ids"], ["BLLD2019-ART21-WAGES"]
        )
        self.assertEqual(result["findings"][0]["risk_level"], "review")
        self.assertEqual(len(result["findings"]), 1)
        self.assertNotIn("compliant", result["summary"].lower())

    @patch("app.urlopen")
    @patch.dict("os.environ", {"LEGAL_AI_API_KEY": "test-key"}, clear=False)
    def test_ai_review_strips_cjk_text_from_provider_output(self, mocked_urlopen):
        class Response:
            def __enter__(self):
                return self

            def __exit__(self, *_):
                return False

            def read(self):
                return json.dumps(
                    {
                        "choices": [
                            {
                                "message": {
                                    "content": json.dumps(
                                        {
                                            "summary": "Sàng lọc sơ bộ 合同风险.",
                                            "findings": [
                                                {
                                                    "title": "Cần kiểm tra nguồn lực 支持资源",
                                                    "issue": "Theo Điều 4.2.3, Bên B sẽ cung cấp哪些具体资源（培训小时数、样品数量、广告预算等） để hỗ trợ Bên A?",
                                                    "risk_level": "review",
                                                    "legal_basis_ids": ["BLLD2019-ART21-WAGES"],
                                                    "missing_facts": ["资源 hỗ trợ"],
                                                    "suggested_revision": "Bổ sung 清单资源 vào phụ lục.",
                                                }
                                            ],
                                            "clarifying_questions": [
                                                {
                                                    "question": "Bên B sẽ cung cấp哪些具体资源 để hỗ trợ Bên A?",
                                                    "clause_ref": "第4.2.3 Điều",
                                                    "why_important": "Thiếu 资源 hỗ trợ.",
                                                }
                                            ],
                                        },
                                        ensure_ascii=False,
                                    )
                                }
                            }
                        ]
                    },
                    ensure_ascii=False,
                ).encode()

        mocked_urlopen.return_value = Response()
        result = review_with_ai(
            "Công việc: Kế toán",
            screen_text("Công việc: Kế toán"),
            consent=True,
        )

        dumped = json.dumps(result, ensure_ascii=False)
        self.assertNotRegex(dumped, r"[\u3400-\u9fff]")
        self.assertIn("Theo Điều 4.2.3", result["findings"][0]["issue"])
        self.assertIn("Bên B sẽ cung cấp", result["clarifying_questions"][0]["question"])

    @patch("app.urlopen")
    @patch.dict("os.environ", {"LEGAL_AI_API_KEY": "test-key"}, clear=False)
    def test_ai_review_drops_placeholder_only_findings_and_questions(self, mocked_urlopen):
        class Response:
            def __enter__(self):
                return self

            def __exit__(self, *_):
                return False

            def read(self):
                return json.dumps(
                    {
                        "choices": [
                            {
                                "message": {
                                    "content": json.dumps(
                                        {
                                            "summary": "Sàng lọc sơ bộ.",
                                            "findings": [
                                                {
                                                    "title": "Cần kiểm tra: Sử dụng placeholder không nhất quán",
                                                    "issue": "Trong hợp đồng xuất hiện {{PERSON_NAME_1}}, {{ADDRESS_1}}, {{TAX_ID_1}} chưa thay thế bằng thông tin thực.",
                                                    "risk_level": "review",
                                                    "legal_basis_ids": ["BLLD2019-ART21-WAGES"],
                                                    "missing_facts": ["{{PERSON_NAME_1}}", "{{ADDRESS_1}}"],
                                                    "suggested_revision": "Thay tất cả placeholder bằng thông tin thực tế.",
                                                },
                                                {
                                                    "title": "Cần kiểm tra mức lương",
                                                    "issue": "Mức lương chưa rõ.",
                                                    "risk_level": "review",
                                                    "legal_basis_ids": ["BLLD2019-ART21-WAGES"],
                                                    "missing_facts": ["Mức lương"],
                                                    "suggested_revision": "Ghi rõ mức lương và thời hạn trả lương.",
                                                },
                                            ],
                                            "clarifying_questions": [
                                                {
                                                    "question": "Vui lòng cung cấp thông tin đầy đủ để thay thế {{PERSON_NAME_1}} và {{ADDRESS_1}}.",
                                                    "clause_ref": "Điều chứa placeholder",
                                                    "why_important": "Tránh mơ hồ về danh tính bên ký.",
                                                },
                                                {
                                                    "question": "Mức lương cụ thể là bao nhiêu?",
                                                    "clause_ref": "Điều 3.1",
                                                    "why_important": "Rủi ro tài chính trực tiếp.",
                                                },
                                            ],
                                        },
                                        ensure_ascii=False,
                                    )
                                }
                            }
                        ]
                    },
                    ensure_ascii=False,
                ).encode()

        mocked_urlopen.return_value = Response()
        result = review_with_ai(
            "Người lao động: {{PERSON_NAME_1}}. Công việc: Kế toán.",
            screen_text("Công việc: Kế toán"),
            consent=True,
        )

        dumped = json.dumps(result, ensure_ascii=False)
        self.assertNotIn("placeholder", dumped.casefold())
        self.assertNotIn("{{PERSON_NAME_1}}", dumped)
        self.assertEqual(len(result["findings"]), 1)
        self.assertIn("mức lương", result["findings"][0]["title"].casefold())
        self.assertEqual(len(result["clarifying_questions"]), 1)
        self.assertIn("mức lương", result["clarifying_questions"][0]["question"].casefold())

    def test_ai_review_persists_clarifying_questions(self):
        class Response:
            def __enter__(self):
                return self

            def __exit__(self, *_):
                return False

            def read(self):
                return json.dumps(
                    {
                        "choices": [
                            {
                                "message": {
                                    "content": json.dumps(
                                        {
                                            "summary": "Sàng lọc sơ bộ.",
                                            "findings": [
                                                {
                                                    "title": "Cần kiểm tra lương",
                                                    "issue": "Thiếu mức lương rõ ràng.",
                                                    "risk_level": "review",
                                                    "legal_basis_ids": [
                                                        "BLLD2019-ART21-WAGES"
                                                    ],
                                                    "missing_facts": ["Mức lương"],
                                                    "suggested_revision": "Ghi rõ mức lương, hình thức trả lương và thời hạn trả lương.",
                                                }
                                            ],
                                            "clarifying_questions": [
                                                {
                                                    "question": "Mức lương cụ thể là bao nhiêu?",
                                                    "clause_ref": "Điều 3.1",
                                                    "why_important": "Rủi ro tài chính trực tiếp.",
                                                },
                                                {
                                                    "question": "",
                                                    "clause_ref": "",
                                                    "why_important": "",
                                                },
                                            ],
                                        },
                                        ensure_ascii=False,
                                    )
                                }
                            }
                        ]
                    },
                    ensure_ascii=False,
                ).encode()

        with patch("app.urlopen", return_value=Response()):
            result = review_with_ai(
                "Công việc: Kế toán",
                screen_text("Công việc: Kế toán"),
                consent=True,
            )

        self.assertEqual(len(result["clarifying_questions"]), 1)
        self.assertIn("Ghi rõ mức lương", result["findings"][0]["suggested_revision"])
        question = result["clarifying_questions"][0]
        self.assertIn("Mức lương", question["question"])
        self.assertIn(question["clause_ref"], "Điều 3.1")

        html = render_report(
            screen_text("Công việc: Kế toán"),
            "contract.txt",
            ai_review=result,
        )
        self.assertIn("Câu hỏi cần làm rõ", html)
        self.assertIn("Mức lương cụ thể", html)
        self.assertIn("Đề xuất chỉnh sửa", html)
        self.assertIn("Ghi rõ mức lương", html)

    @patch("app.urlopen")
    @patch.dict("os.environ", {"LEGAL_AI_API_KEY": "test-key"}, clear=False)
    def test_ai_prompt_asks_for_suggested_revisions_and_ignores_pii_placeholders(self, mocked_urlopen):
        captured: dict[str, str] = {}

        class Response:
            def __enter__(self):
                return self

            def __exit__(self, *_):
                return False

            def read(self):
                return json.dumps(
                    {"choices": [{"message": {"content": json.dumps({"summary": "OK", "findings": [], "clarifying_questions": []}, ensure_ascii=False)}}]},
                    ensure_ascii=False,
                ).encode()

        def fake_urlopen(request, timeout):
            captured["body"] = request.data.decode()
            return Response()

        mocked_urlopen.side_effect = fake_urlopen
        review_with_ai(
            "Người lao động: Nguyễn Văn A. Công việc: Kế toán.",
            screen_text("Công việc: Kế toán"),
            consent=True,
        )

        body = captured["body"]
        self.assertIn("suggested_revision", body)
        self.assertIn("Đề xuất chỉnh sửa", body)
        self.assertIn("Không coi placeholder PII", body)
        self.assertIn("{{PERSON_NAME_", body)
        self.assertIn("Chỉ viết tiếng Việt", body)
        self.assertIn("không dùng tiếng Trung", body)

    def test_ai_review_without_clarifying_questions_renders_no_section(self):
        result = {
            "summary": "Sàng lọc sơ bộ.",
            "findings": [],
            "clarifying_questions": [],
        }
        html = render_report(
            screen_text("Công việc: Kế toán"),
            "contract.txt",
            ai_review=result,
        )
        self.assertNotIn("Câu hỏi cần làm rõ", html)


    def test_ai_review_masks_pii_before_sending_to_provider(self):
        captured: dict[str, str] = {}

        class Response:
            def __enter__(self):
                return self

            def __exit__(self, *_):
                return False

            def read(self):
                return json.dumps(
                    {"choices": [{"message": {"content": json.dumps({"summary": "OK", "findings": [], "clarifying_questions": []}, ensure_ascii=False)}}]},
                    ensure_ascii=False,
                ).encode()

        def fake_urlopen(request, timeout):
            captured["body"] = request.data.decode()
            return Response()

        contract = "Người lao động: Nguyễn Văn A, SĐT 0901234567, email nguyenvana@gmail.com, CCCD 079203001234."
        with patch("app.urlopen", side_effect=fake_urlopen):
            review_with_ai(contract, screen_text("Công việc: Kế toán"), consent=True)

        body = captured["body"]
        self.assertNotIn("0901234567", body)
        self.assertNotIn("nguyenvana@gmail.com", body)
        self.assertNotIn("079203001234", body)
        self.assertIn("PHONE_NUMBER", body)
        self.assertIn("EMAIL_ADDRESS", body)
        self.assertIn("VN_CCCD", body)

    def test_chat_session_renders_chat_form_after_ai_review(self):
        CHAT_SESSIONS.clear()
        result = screen_text("Công việc: Kế toán")
        ai_review = {"summary": "OK", "findings": [], "clarifying_questions": []}
        session_id = create_chat_session("contract.txt", "Công việc: Kế toán", result, ai_review)

        html = render_report(result, "contract.txt", ai_review=ai_review, session_id=session_id)

        self.assertIn("Hỏi tiếp về hợp đồng", html)
        self.assertIn(session_id, html)
        self.assertIn("← Kiểm tra hợp đồng khác", html)

    @patch("app.urlopen")
    def test_chat_answer_renders_inside_existing_report_with_history(self, mocked_urlopen):
        CHAT_SESSIONS.clear()
        result = screen_text("Công việc: Kế toán")
        ai_review = {"summary": "OK", "findings": [], "clarifying_questions": []}
        session_id = create_chat_session("contract.txt", "Công việc: Kế toán", result, ai_review)

        class Response:
            def __enter__(self):
                return self

            def __exit__(self, *_):
                return False

            def read(self):
                return json.dumps({"choices": [{"message": {"content": "Cần kiểm tra lương."}}]}, ensure_ascii=False).encode()

        mocked_urlopen.return_value = Response()
        answer = chat_with_ai(session_id, "Lương ổn chưa?")
        html = render_chat_page(session_id)

        self.assertIn("Kết quả rà soát sơ bộ", html)
        self.assertIn("Lương ổn chưa?", html)
        self.assertIn(answer, html)
        self.assertIn("Hỏi tiếp về hợp đồng", html)
        self.assertIn("id='chat-panel'", html)
        self.assertIn("class='chat-log'", html)
        self.assertIn("action='/chat#chat-panel'", html)
        self.assertIn("scrollTop = chatLog.scrollHeight", html)

    @patch("app.urlopen")
    @patch.dict("os.environ", {"LEGAL_AI_API_KEY": "test-key"}, clear=False)
    def test_chat_masks_user_question_before_sending_to_provider(self, mocked_urlopen):
        CHAT_SESSIONS.clear()
        result = screen_text("Công việc: Kế toán")
        session_id = create_chat_session(
            "contract.txt",
            "Người lao động SĐT 0901234567. Công việc: Kế toán.",
            result,
            {"summary": "OK", "findings": [], "clarifying_questions": []},
        )
        captured: dict[str, str] = {}

        class Response:
            def __enter__(self):
                return self

            def __exit__(self, *_):
                return False

            def read(self):
                return json.dumps({"choices": [{"message": {"content": "Cần kiểm tra mức lương."}}]}, ensure_ascii=False).encode()

        def fake_urlopen(request, timeout):
            captured["body"] = request.data.decode()
            return Response()

        mocked_urlopen.side_effect = fake_urlopen
        answer = chat_with_ai(session_id, "SĐT 0901234567 có cần ghi không?")

        self.assertIn("Cần kiểm tra", answer)
        self.assertNotIn("0901234567", captured["body"])
        self.assertIn("PHONE_NUMBER", captured["body"])
        self.assertIn("Chỉ viết tiếng Việt", captured["body"])


if __name__ == "__main__":
    unittest.main()
